from datetime import datetime, time, timedelta, date
from io import BytesIO
from typing import Any, Dict, List, Optional, Tuple
from openpyxl import load_workbook

# Import kalkulator SLA 3
from app.sla.sla3.calculator import SLA3Calculator


def _raw(value):
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, time):
        return value.isoformat()
    return str(value)


def _as_date(value) -> Optional[date]:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d %b %Y", "%d %B %Y", "%d/%m/%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    try:
        return datetime.fromisoformat(text).date()
    except Exception:
        return None


def _as_time(value) -> Optional[time]:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.time()
    if isinstance(value, time):
        return value
    if isinstance(value, (int, float)) and 0 <= float(value) < 1:
        return (datetime.min + timedelta(seconds=round(float(value) * 86400))).time()
    text = str(value).strip()
    for fmt in ("%H:%M:%S", "%H:%M", "%I:%M:%S %p", "%I:%M %p"):
        try:
            return datetime.strptime(text, fmt).time()
        except ValueError:
            pass
    return None


def _combine(d, t):
    dd, tt = _as_date(d), _as_time(t)
    return datetime.combine(dd, tt) if dd and tt else None


def _norm_header(value: Any) -> str:
    return " ".join(str(value or "").strip().lower().replace("\n", " ").split())


CANONICAL_HEADERS = {
    "sumber": ["Sumber"],
    "status_registrasi": ["Status Registrasi"],
    "status_inisiasi": ["Status Inisiasi"],
    "posisi_data": ["Posisi Data"],
    "nomor_registrasi": ["Nomor Registrasi"],
    "nama_dokumen": ["Nama Dokumen / Hal", "Nama Dokumen", "Nama Dokumen/Hal"],
    "unit": ["Unit", "Unit Kerja"],
    "format_arsip": ["Format Arsip"],
    "document_type": ["Document Type"],
    "tanggal_registrasi": ["Tanggal Registrasi"],
    "jam_registrasi": ["Jam Registrasi"],
    "tanggal_verifikasi_uf": ["Tanggal Verifikasi UF"],
    "jam_verifikasi_uf": ["Jam Verifikasi UF"],
    "tanggal_verifikasi_uu": ["Tanggal Verifikasi UU"],
    "jam_verifikasi_uu": ["Jam Verifikasi UU"],
    "tanggal_verifikasi_arsip": ["Tanggal Verifikasi Arsip"],
    "jam_verifikasi_arsip": ["Jam Verifikasi Arsip"],
    "tanggal_registrasi_2": ["Tanggal Registrasi 2"],
    "jam_registrasi_2": ["Jam Registrasi 2"],
    "tanggal_verifikasi_arsip_2": ["Tanggal Verifikasi Arsip 2"],
    "jam_verifikasi_arsip_2": ["Jam Verifikasi Arsip 2"],
    "tanggal_penjadwalan_inaktif": ["Tanggal Penjadwalan Arsip Inaktif"],
    "jam_penjadwalan_inaktif": ["Jam Penjadwalan Arsip Inaktif"],
    "tanggal_penolakan_jadwal_penjemputan": ["Tanggal Penolakan Jadwal Penjemputan"],
    "jam_penolakan_jadwal_penjemputan": ["Jam Penolakan Jadwal Penjemputan"],
    "tanggal_gagal_penjemputan": ["Tanggal Gagal Penjemputan"],
    "jam_gagal_penjemputan": ["Jam Gagal Penjemputan"],
    "tanggal_jadwal_penjemputan": ["Tanggal Jadwal Penjemputan"],
    "jam_jadwal_penjemputan": ["Jam Jadwal Penjemputan"],
    "tanggal_penjemputan_dokumen": ["Tanggal Penjemputan Dokumen"],
    "jam_penjemputan_dokumen": ["Jam Penjemputan Dokumen"],
    "tanggal_penjadwalan_inaktif_2": ["Tanggal Penjadwalan Arsip Inaktif 2"],
    "jam_penjadwalan_inaktif_2": ["Jam Penjadwalan Arsip Inaktif 2"],
    "tanggal_runner_record_center": ["Tanggal Runner sampai di Record Center", "Tanggal Runner sampai di RC"],
    "jam_runner_record_center": ["Jam Runner sampai di Record Center", "Jam Runner sampai di RC"],
    "tanggal_registrasi_arsip_generate_barcode": ["Tanggal Registrasi Arsip di Generate Barcode"],
    "jam_registrasi_arsip_generate_barcode": ["Jam Registrasi Arsip di Generate Barcode"],
}


def _find_header(ws):
    aliases = {_norm_header(a): key for key, vals in CANONICAL_HEADERS.items() for a in vals}
    best = None
    for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 40), values_only=True), start=1):
        mapping = {}
        for idx, value in enumerate(row):
            n = _norm_header(value)
            if n in aliases:
                mapping[aliases[n]] = idx
        score = len(mapping)
        if best is None or score > best[0]:
            best = (score, row_idx, mapping)
    if not best or best[0] < 10:
        raise ValueError("Header raw data tidak dikenali. Pastikan nama kolom SLA 1 tersedia.")
    return best[1], best[2]


def _last_weekday(year: int, month: int, nth: int):
    if month == 12:
        cur = date(year + 1, 1, 1) - timedelta(days=1)
    else:
        cur = date(year, month + 1, 1) - timedelta(days=1)
    count = 0
    while True:
        if cur.weekday() < 5:
            count += 1
            if count == nth:
                return cur
        cur -= timedelta(days=1)


def build_period_context(reg_dates: List[date], holidays: List[date], cutoff: time = time(10, 0)) -> Dict[str, Any]:
    if not reg_dates:
        raise ValueError("Tidak ditemukan Tanggal Registrasi yang valid.")
    latest = max(reg_dates)
    y, m = latest.year, latest.month
    prev_y, prev_m = (y - 1, 12) if m == 1 else (y, m - 1)
    current_last_working_day_1 = _last_weekday(y, m, 1)
    current_last_working_day_2 = _last_weekday(y, m, 2)
    current_last_working_day_3 = _last_weekday(y, m, 3)
    previous_last_working_day_1 = _last_weekday(prev_y, prev_m, 1)
    previous_last_working_day_2 = _last_weekday(prev_y, prev_m, 2)
    previous_last_working_day_3 = _last_weekday(prev_y, prev_m, 3)

    return {
        "current_last_working_day_1": current_last_working_day_1,
        "current_last_working_day_2": current_last_working_day_2,
        "current_last_working_day_3": current_last_working_day_3,
        "previous_last_working_day_1": previous_last_working_day_1,
        "previous_last_working_day_2": previous_last_working_day_2,
        "previous_last_working_day_3": previous_last_working_day_3,
        "period_reference_cells": {
            "H5": current_last_working_day_1,
            "H6": current_last_working_day_2,
            "H7": current_last_working_day_3,
            "H8": previous_last_working_day_1,
            "H9": previous_last_working_day_2,
            "H10": previous_last_working_day_3,
        },
        "h5": current_last_working_day_1,
        "h6": current_last_working_day_2,
        "h8": previous_last_working_day_1,
        "h9": previous_last_working_day_2,
        "cutoff": cutoff,
        "holidays": holidays,
        "period_label": f"{y:04d}-{m:02d}",
    }


def parse_registrasi_sheet(file_bytes: bytes, holidays: Optional[List[date]] = None) -> Tuple[List[Dict[str, Any]], Dict[str, Any]]:
    wb = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    try:
        candidates = []
        for ws in wb.worksheets:
            try:
                header_row, mapping = _find_header(ws)
                candidates.append((len(mapping), ws, header_row, mapping))
            except ValueError:
                continue
        if not candidates:
            raise ValueError("Tidak menemukan sheet raw registrasi dengan header SLA 1.")
        _, ws, header_row, mapping = max(candidates, key=lambda x: x[0])

        rows = ws.iter_rows(min_row=header_row + 1, values_only=True)
        required = ["sumber", "status_registrasi", "format_arsip", "document_type", "tanggal_registrasi", "tanggal_verifikasi_arsip"]
        missing = [x for x in required if x not in mapping]
        if missing:
            raise ValueError("Kolom wajib raw data tidak ditemukan: " + ", ".join(missing))

        def val(row, key):
            idx = mapping.get(key)
            return row[idx] if idx is not None and idx < len(row) else None

        records = []
        reg_dates = []
        for row in rows:
            nom = val(row, "nomor_registrasi")
            name = val(row, "nama_dokumen")
            reg_d = val(row, "tanggal_registrasi")
            if all(v is None or str(v).strip() == "" for v in (nom, name, reg_d)):
                continue
            rd = _as_date(reg_d)
            if rd:
                reg_dates.append(rd)
            rt = val(row, "jam_registrasi")
            vuf_d, vuf_t = val(row, "tanggal_verifikasi_uf"), val(row, "jam_verifikasi_uf")
            vuu_d, vuu_t = val(row, "tanggal_verifikasi_uu"), val(row, "jam_verifikasi_uu")
            va_d, va_t = val(row, "tanggal_verifikasi_arsip"), val(row, "jam_verifikasi_arsip")
            r2_d, r2_t = val(row, "tanggal_registrasi_2"), val(row, "jam_registrasi_2")
            va2_d, va2_t = val(row, "tanggal_verifikasi_arsip_2"), val(row, "jam_verifikasi_arsip_2")
            ti_d, ti_t = val(row, "tanggal_penjadwalan_inaktif"), val(row, "jam_penjadwalan_inaktif")
            pf_d, pf_t = val(row, "tanggal_penolakan_jadwal_penjemputan"), val(row, "jam_penolakan_jadwal_penjemputan")
            gf_d, gf_t = val(row, "tanggal_gagal_penjemputan"), val(row, "jam_gagal_penjemputan")
            jp_d, jp_t = val(row, "tanggal_jadwal_penjemputan"), val(row, "jam_jadwal_penjemputan")
            pd_d, pd_t = val(row, "tanggal_penjemputan_dokumen"), val(row, "jam_penjemputan_dokumen")
            ti2_d, ti2_t = val(row, "tanggal_penjadwalan_inaktif_2"), val(row, "jam_penjadwalan_inaktif_2")
            rr_d, rr_t = val(row, "tanggal_runner_record_center"), val(row, "jam_runner_record_center")
            barcode_d, barcode_t = val(row, "tanggal_registrasi_arsip_generate_barcode"), val(row, "jam_registrasi_arsip_generate_barcode")
            
            records.append({
                "sumber": _raw(val(row, "sumber")), "status_registrasi": _raw(val(row, "status_registrasi")),
                "status_inisiasi": _raw(val(row, "status_inisiasi")), "posisi_data": _raw(val(row, "posisi_data")),
                "nomor_registrasi": _raw(nom), "nama_dokumen": _raw(name), "unit": _raw(val(row, "unit")) or "Tanpa Unit",
                "format_arsip": _raw(val(row, "format_arsip")), "document_type": _raw(val(row, "document_type")),
                "tanggal_registrasi": rd, "jam_registrasi": _as_time(rt),
                "tanggal_verifikasi_uf": _as_date(vuf_d), "jam_verifikasi_uf": _as_time(vuf_t),
                "tanggal_verifikasi_uu": _as_date(vuu_d), "jam_verifikasi_uu": _as_time(vuu_t),
                "tanggal_verifikasi_arsip": _as_date(va_d), "jam_verifikasi_arsip": _as_time(va_t),
                "tanggal_registrasi_2": _as_date(r2_d), "jam_registrasi_2": _as_time(r2_t),
                "tanggal_verifikasi_arsip_2": _as_date(va2_d), "jam_verifikasi_arsip_2": _as_time(va2_t),
                "tanggal_penjadwalan_inaktif": _as_date(ti_d), "jam_penjadwalan_inaktif": _as_time(ti_t),
                "tanggal_penolakan_jadwal_penjemputan": _as_date(pf_d), "jam_penolakan_jadwal_penjemputan": _as_time(pf_t),
                "tanggal_gagal_penjemputan": _as_date(gf_d), "jam_gagal_penjemputan": _as_time(gf_t),
                "tanggal_jadwal_penjemputan": _as_date(jp_d), "jam_jadwal_penjemputan": _as_time(jp_t),
                "tanggal_penjemputan_dokumen": _as_date(pd_d), "jam_penjemputan_dokumen": _as_time(pd_t),
                "tanggal_penjadwalan_inaktif_2": _as_date(ti2_d), "jam_penjadwalan_inaktif_2": _as_time(ti2_t),
                "tanggal_runner_record_center": _as_date(rr_d), "jam_runner_record_center": _as_time(rr_t),
                "tanggal_registrasi_arsip_generate_barcode": _as_date(barcode_d), "jam_registrasi_arsip_generate_barcode": _as_time(barcode_t),
                "registration_timestamp": _combine(reg_d, rt), "verification_uf_timestamp": _combine(vuf_d, vuf_t),
                "verification_uu_timestamp": _combine(vuu_d, vuu_t), "verification_timestamp": _combine(va_d, va_t),
                "registration_2_timestamp": _combine(r2_d, r2_t), "verification_2_timestamp": _combine(va2_d, va2_t),
                "penjadwalan_inaktif_timestamp": _combine(ti_d, ti_t),
                "penolakan_jadwal_penjemputan_timestamp": _combine(pf_d, pf_t),
                "gagal_penjemputan_timestamp": _combine(gf_d, gf_t),
                "penjadwalan_inaktif_2_timestamp": _combine(ti2_d, ti2_t),
            })
        
        # Build period context (sekarang kita punya H5, H6, H8, H9)
        context = build_period_context(reg_dates, holidays or [])

        # --- EVALUASI SLA 3 PADA MASING-MASING RECORD ---
        for rec in records:
            sla3_eval = SLA3Calculator.evaluate_vendor(rec, context)
            
            # Titipkan hasil evaluasi di dalam dictionary untuk disimpan di rute API
            rec["sla3_eval"] = sla3_eval
            rec["sla3_dashboard_status"] = SLA3Calculator.to_dashboard_status(sla3_eval["result"])

        return records, context
    finally:
        wb.close()