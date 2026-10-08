from datetime import date, datetime, time, timedelta
import io
from typing import Optional

import pandas as pd
from fastapi import APIRouter, Depends, File, HTTPException, UploadFile, Response
from sqlalchemy import func, case, or_, and_
from sqlalchemy.orm import Session
from openpyxl import load_workbook

from app.database import get_db
from app.models.sla_record import SLARecord
from app.models.holiday import HolidayCalendar
from app.services.excel_parser import parse_registrasi_sheet
from app.sla.sla1.calculator import SLA1Calculator
from app.sla.sla2.calculator import SLA2Calculator
from app.sla.sla3.calculator import SLA3Calculator
from app.api.dashboard_metrics import build_dashboard

router = APIRouter()
SLA_CODE = "SLA 1"


def _apply_filters(query, start_date=None, end_date=None, unit=None, status=None):
    if start_date:
        query = query.filter(SLARecord.registration_timestamp >= start_date)
    if end_date:
        query = query.filter(SLARecord.registration_timestamp <= end_date + " 23:59:59")
    if unit:
        query = query.filter(SLARecord.unit.ilike(f"%{unit}%"))
    if status:
        query = query.filter(SLARecord.status == status)
    return query


def _active_holidays(db: Session, year: int, version: Optional[str] = None):
    q = db.query(HolidayCalendar).filter(HolidayCalendar.year == year)
    if version:
        q = q.filter(HolidayCalendar.version == version)
    else:
        latest = q.order_by(HolidayCalendar.id.desc()).first()
        if latest:
            q = q.filter(HolidayCalendar.version == latest.version)
    return [x.holiday_date for x in q.order_by(HolidayCalendar.holiday_date).all()]


def _latest_version(db: Session, year: int):
    row = db.query(HolidayCalendar.version).filter(HolidayCalendar.year == year).order_by(HolidayCalendar.id.desc()).first()
    return row[0] if row else None


def _default_period_from_records(records):
    dates = [r.registration_timestamp.date() for r in records if r.registration_timestamp]
    return max(dates).strftime("%Y-%m") if dates else None


@router.get('/health')
def health():
    return {"status": "ok", "scope": "SLA 1 + SLA 2 + SLA 3"}


@router.post('/import')
async def import_excel(file: UploadFile = File(...), db: Session = Depends(get_db)):
    contents = await file.read()
    try:
        # Holiday calendar dibaca setelah kita mengetahui tahun periode data.
        raw_records, temp_context = parse_registrasi_sheet(contents, holidays=[])
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Gagal membaca raw Excel: {e}")

    if not raw_records:
        raise HTTPException(status_code=400, detail="Tidak ada baris raw data yang dapat diproses.")

    period = temp_context["period_label"]
    year = int(period[:4])
    version = _latest_version(db, year)
    holidays = _active_holidays(db, year, version) if version else []
    context = dict(temp_context)
    context["holidays"] = holidays
    context["holiday_calendar_version"] = version or "UNCONFIGURED"

    # Import baru menggantikan data SLA 1 lama pada periode yang sama.
    start_period = datetime(year, int(period[5:7]), 1)
    end_period = datetime(year + 1, 1, 1) if int(period[5:7]) == 12 else datetime(year, int(period[5:7]) + 1, 1)
    # Import baru menggantikan dataset pada periode yang sama. Gunakan
    # reporting_period sebagai identitas utama; fallback timestamp dipakai
    # untuk baris lama yang belum mempunyai reporting_period.
    db.query(SLARecord).filter(
        SLARecord.sla_code == SLA_CODE,
        or_(
            SLARecord.reporting_period == period,
            and_(
                SLARecord.reporting_period.is_(None),
                SLARecord.registration_timestamp >= start_period,
                SLARecord.registration_timestamp < end_period,
            ),
        ),
    ).delete(synchronize_session=False)

    mappings = []
    sla1_counts = {"Yes": 0, "No": 0, "N/A": 0}
    sla2_counts = {"Yes": 0, "No": 0, "N/A": 0}
    sla2a_counts = {"Yes": 0, "No": 0, "N/A": 0}
    sla2b_counts = {"Yes": 0, "No": 0, "N/A": 0}
    sla3_counts = {"Yes": 0, "No": 0, "N/A": 0}
    sla3a_counts = {"Yes": 0, "No": 0, "N/A": 0}
    sla3b_counts = {"Yes": 0, "No": 0, "N/A": 0}

    for rec in raw_records:
        # SLA 1 tetap dihitung seperti sebelumnya.
        evaluation1 = SLA1Calculator.evaluate_vendor(rec, context)
        result1 = evaluation1["result"]

        # SLA 2A dan SLA 2B dihitung paralel sesuai dua kolom formula Excel.
        evaluation2 = SLA2Calculator.evaluate_vendor(rec, context)
        result2 = evaluation2["result"]

        # SLA 3A + SLA 3B + final fallback/manual logic.
        evaluation3 = SLA3Calculator.evaluate_vendor(rec, context)
        result3 = evaluation3["result"]

        sla1_counts[result1] += 1
        sla2_counts[result2] += 1
        sla2a_counts[evaluation2["sla_2a_result"]] += 1
        sla2b_counts[evaluation2["sla_2b_result"]] += 1
        sla3_counts[result3] += 1
        sla3a_counts[evaluation3.get("sla_3a_result", "N/A")] += 1
        sla3b_counts[evaluation3.get("sla_3b_result", "N/A")] += 1

        mappings.append({
            "nomor_registrasi": rec.get("nomor_registrasi"), "nama_dokumen": rec.get("nama_dokumen"), "unit": rec.get("unit"),
            "sumber": rec.get("sumber"), "status_registrasi": rec.get("status_registrasi"), "status_inisiasi": rec.get("status_inisiasi"),
            "posisi_data": rec.get("posisi_data"), "format_arsip": rec.get("format_arsip"), "document_type": rec.get("document_type"),
            "tanggal_registrasi_raw": str(rec.get("tanggal_registrasi") or ""), "jam_registrasi_raw": str(rec.get("jam_registrasi") or ""),
            "tanggal_verifikasi_uf_raw": str(rec.get("tanggal_verifikasi_uf") or ""), "jam_verifikasi_uf_raw": str(rec.get("jam_verifikasi_uf") or ""),
            "tanggal_verifikasi_uu_raw": str(rec.get("tanggal_verifikasi_uu") or ""), "jam_verifikasi_uu_raw": str(rec.get("jam_verifikasi_uu") or ""),
            "tanggal_verifikasi_arsip_raw": str(rec.get("tanggal_verifikasi_arsip") or ""), "jam_verifikasi_arsip_raw": str(rec.get("jam_verifikasi_arsip") or ""),
            "tanggal_registrasi_2_raw": str(rec.get("tanggal_registrasi_2") or ""), "jam_registrasi_2_raw": str(rec.get("jam_registrasi_2") or ""),
            "tanggal_verifikasi_arsip_2_raw": str(rec.get("tanggal_verifikasi_arsip_2") or ""), "jam_verifikasi_arsip_2_raw": str(rec.get("jam_verifikasi_arsip_2") or ""),
            "tanggal_penjadwalan_inaktif_raw": str(rec.get("tanggal_penjadwalan_inaktif") or ""), "jam_penjadwalan_inaktif_raw": str(rec.get("jam_penjadwalan_inaktif") or ""),
            "tanggal_penolakan_jadwal_penjemputan_raw": str(rec.get("tanggal_penolakan_jadwal_penjemputan") or ""), "jam_penolakan_jadwal_penjemputan_raw": str(rec.get("jam_penolakan_jadwal_penjemputan") or ""),
            "tanggal_gagal_penjemputan_raw": str(rec.get("tanggal_gagal_penjemputan") or ""), "jam_gagal_penjemputan_raw": str(rec.get("jam_gagal_penjemputan") or ""),
            "tanggal_jadwal_penjemputan_raw": str(rec.get("tanggal_jadwal_penjemputan") or ""), "jam_jadwal_penjemputan_raw": str(rec.get("jam_jadwal_penjemputan") or ""),
            "tanggal_penjemputan_dokumen_raw": str(rec.get("tanggal_penjemputan_dokumen") or ""), "jam_penjemputan_dokumen_raw": str(rec.get("jam_penjemputan_dokumen") or ""),
            "tanggal_penjadwalan_inaktif_2_raw": str(rec.get("tanggal_penjadwalan_inaktif_2") or ""), "jam_penjadwalan_inaktif_2_raw": str(rec.get("jam_penjadwalan_inaktif_2") or ""),
            "registration_timestamp": rec.get("registration_timestamp"), "verification_uf_timestamp": rec.get("verification_uf_timestamp"),
            "verification_uu_timestamp": rec.get("verification_uu_timestamp"), "verification_timestamp": rec.get("verification_timestamp"),
            "registration_2_timestamp": rec.get("registration_2_timestamp"), "verification_2_timestamp": rec.get("verification_2_timestamp"),
            "sla_code": SLA_CODE, "status": SLA1Calculator.to_dashboard_status(result1),
            "incomplete_reason": evaluation1.get("reason"), "is_denominator": 1,
            "sla1a_result": evaluation1.get("sla_1a_result"), "sla1b_result": evaluation1.get("sla_1b_result"),
            "sla1_decision_source": evaluation1.get("decision_source"), "sla1_manual_rule": evaluation1.get("manual_rule"),
            "sla1_working_days": evaluation1.get("working_days"),
            "sla2a_result": evaluation2.get("sla_2a_result"), "sla2b_result": evaluation2.get("sla_2b_result"),
            "sla2_final_result": result2, "sla2_manual_result": evaluation2.get("manual_result"),
            "sla2_decision_source": evaluation2.get("decision_source"), "sla2_manual_rule": evaluation2.get("manual_rule"),
            "sla2_working_days": evaluation2.get("working_days"),
            "sla2a_working_days": evaluation2.get("sla_2a_working_days"), "sla2b_working_days": evaluation2.get("sla_2b_working_days"),
            "sla2_flow_stage": evaluation2.get("flow_stage"), "sla2_reason": evaluation2.get("reason"),
            "sla3a_result": evaluation3.get("sla_3a_result"), "sla3b_result": evaluation3.get("sla_3b_result"),
            "sla3_final_result": result3, "sla3_manual_result": evaluation3.get("manual_result"),
            "sla3_decision_source": evaluation3.get("decision_source"), "sla3_manual_rule": evaluation3.get("manual_rule"),
            "sla3_working_days": evaluation3.get("working_days"),
            "sla3a_working_days": evaluation3.get("sla_3a_working_days"), "sla3b_working_days": evaluation3.get("sla_3b_working_days"),
            "sla3_flow_stage": evaluation3.get("flow_stage"), "sla3_reason": evaluation3.get("reason"),
            "holiday_calendar_version": context["holiday_calendar_version"],
            "flow_stage": evaluation1.get("flow_stage"),
            "reporting_period": period,
        })

    db.bulk_insert_mappings(SLARecord, mappings)
    db.commit()
    return {
        "message": f"Berhasil memproses {len(raw_records)} baris raw data dengan perhitungan SLA 1, SLA 2, dan SLA 3 periode {period}.",
        "records": len(raw_records), "period": period, "holiday_calendar_version": context["holiday_calendar_version"],
        "holidays_used": len(holidays),
        "sla1": sla1_counts,
        "sla2": {"final": sla2_counts, "sla2a": sla2a_counts, "sla2b": sla2b_counts},
        "sla3": {"final": sla3_counts, "sla3a": sla3a_counts, "sla3b": sla3b_counts},
    }


@router.get('/monitoring/dashboard')
def monitoring_dashboard(
    sla: str = "ALL",
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    table_unit: Optional[str] = None,
    trend_unit: Optional[str] = None,
    db: Session = Depends(get_db),
):
    """Single dashboard payload for SLA 1/SLA 2, unit table, ranking and daily trend."""
    return build_dashboard(
        db,
        start_date=start_date,
        end_date=end_date,
        sla=sla,
        table_unit=table_unit,
        trend_unit=trend_unit,
    )


@router.get('/monitoring/summary')
def summary(start_date: Optional[str] = None, end_date: Optional[str] = None, unit: Optional[str] = None, status: Optional[str] = None, db: Session = Depends(get_db)):
    q = _apply_filters(db.query(SLARecord), start_date, end_date, unit, status)
    row = q.with_entities(
        func.count(SLARecord.id).label('total'),
        func.sum(case((SLARecord.is_denominator == 1, 1), else_=0)).label('denominator'),
        func.sum(case((SLARecord.status == 'ON TIME', 1), else_=0)).label('on_time'),
        func.sum(case((SLARecord.status == 'OUT OF DATE', 1), else_=0)).label('out_of_date'),
        func.sum(case((SLARecord.status == 'INCOMPLETE', 1), else_=0)).label('incomplete'),
    ).filter(SLARecord.sla_code == SLA_CODE).first()
    total, den, on_time, late, incomplete = [int(getattr(row, k) or 0) for k in ('total','denominator','on_time','out_of_date','incomplete')]
    pct = round(on_time / den * 100, 2) if den else 0.0

    unit_rows = q.filter(SLARecord.sla_code == SLA_CODE).with_entities(
        SLARecord.unit.label('unit'),
        func.sum(case((SLARecord.is_denominator == 1, 1), else_=0)).label('denominator'),
        func.sum(case((SLARecord.status == 'ON TIME', 1), else_=0)).label('on_time'),
        func.sum(case((SLARecord.status == 'OUT OF DATE', 1), else_=0)).label('late'),
    ).group_by(SLARecord.unit).all()
    units = []
    for r in unit_rows:
        d, o = int(r.denominator or 0), int(r.on_time or 0)
        p = round(o / d * 100, 2) if d else 0.0
        units.append({"unit": r.unit or "Tanpa Unit", "denominator": d, "on_time": o, "late": int(r.late or 0), "percentage": p, "tindak_lanjut": "Pertahankan capaian." if p >= 100 else "Evaluasi berkas yang belum memenuhi SLA."})
    units.sort(key=lambda x: (x['percentage'], x['denominator']), reverse=True)
    top5, bottom5 = units[:5], sorted(units, key=lambda x: (x['percentage'], -x['denominator']))[:5]
    return {
        "sla": "SLA 1", "accumulation": {"total_records": total, "denominator": den, "on_time": on_time, "out_of_date": late, "incomplete": incomplete, "percentage": pct, "ontime_pct": pct, "ood_pct": round(late / den * 100, 2) if den else 0.0, "inc_pct": round(incomplete / den * 100, 2) if den else 0.0},
        "unit_summary": units, "ranking": {"top5": top5, "bottom5": bottom5},
    }


@router.get('/monitoring/trend')
def trend(start_date: Optional[str] = None, end_date: Optional[str] = None, unit: Optional[str] = None, status: Optional[str] = None, db: Session = Depends(get_db)):
    q = _apply_filters(db.query(SLARecord), start_date, end_date, unit, status).filter(SLARecord.sla_code == SLA_CODE, SLARecord.registration_timestamp.isnot(None), SLARecord.is_denominator == 1)
    rows = q.with_entities(func.date(SLARecord.registration_timestamp).label('tanggal'), func.count(SLARecord.id).label('jumlah_arsip'), func.sum(case((SLARecord.status == 'ON TIME', 1), else_=0)).label('jumlah_berhasil')).group_by('tanggal').order_by('tanggal').all()
    return {"points": [{"tanggal": r.tanggal, "jumlah_arsip": int(r.jumlah_arsip or 0), "jumlah_berhasil": int(r.jumlah_berhasil or 0), "persentase": round((r.jumlah_berhasil or 0) / r.jumlah_arsip * 100, 2) if r.jumlah_arsip else None} for r in rows]}


@router.get('/monitoring/traceability')
def traceability(page: int = 1, limit: int = 150, start_date: Optional[str] = None, end_date: Optional[str] = None, unit: Optional[str] = None, status: Optional[str] = None, db: Session = Depends(get_db)):
    q = _apply_filters(db.query(SLARecord).filter(SLARecord.sla_code == SLA_CODE), start_date, end_date, unit, status)
    total = q.count()
    records = q.order_by(SLARecord.id.desc()).offset((page - 1) * limit).limit(limit).all()
    units = sorted([x[0] for x in db.query(SLARecord.unit).filter(SLARecord.sla_code == SLA_CODE).distinct().all() if x[0]])
    return {"total": total, "page": page, "limit": limit, "units": units, "data": records}



def _sla2_status_filter(query, status: Optional[str]):
    if not status:
        return query
    mapping = {"ON TIME": "Yes", "OUT OF DATE": "No", "INCOMPLETE": "N/A"}
    value = mapping.get(status)
    return query.filter(SLARecord.sla2_final_result == value) if value else query


@router.get('/monitoring/sla2/summary')
def sla2_summary(start_date: Optional[str] = None, end_date: Optional[str] = None, unit: Optional[str] = None, status: Optional[str] = None, db: Session = Depends(get_db)):
    q = db.query(SLARecord).filter(SLARecord.sla_code == SLA_CODE)
    if start_date:
        q = q.filter(SLARecord.registration_timestamp >= start_date)
    if end_date:
        q = q.filter(SLARecord.registration_timestamp <= end_date + " 23:59:59")
    if unit:
        q = q.filter(SLARecord.unit.ilike(f"%{unit}%"))
    q = _sla2_status_filter(q, status)
    row = q.with_entities(
        func.count(SLARecord.id).label('total'),
        func.sum(case((SLARecord.sla2_final_result == 'Yes', 1), else_=0)).label('on_time'),
        func.sum(case((SLARecord.sla2_final_result == 'No', 1), else_=0)).label('out_of_date'),
        func.sum(case((SLARecord.sla2_final_result == 'N/A', 1), else_=0)).label('incomplete'),
    ).first()
    total = int(row.total or 0); on_time = int(row.on_time or 0); late = int(row.out_of_date or 0); incomplete = int(row.incomplete or 0)
    denominator = on_time + late
    pct = round(on_time / denominator * 100, 2) if denominator else 0.0

    unit_rows = q.with_entities(
        SLARecord.unit.label('unit'),
        func.sum(case((SLARecord.sla2_final_result == 'Yes', 1), else_=0)).label('on_time'),
        func.sum(case((SLARecord.sla2_final_result == 'No', 1), else_=0)).label('late'),
        func.sum(case((SLARecord.sla2_final_result == 'N/A', 1), else_=0)).label('incomplete'),
    ).group_by(SLARecord.unit).all()
    units = []
    for r in unit_rows:
        o, l = int(r.on_time or 0), int(r.late or 0); d = o + l
        p = round(o / d * 100, 2) if d else 0.0
        units.append({"unit": r.unit or "Tanpa Unit", "denominator": d, "on_time": o, "late": l, "incomplete": int(r.incomplete or 0), "percentage": p, "tindak_lanjut": "Pertahankan capaian." if p >= 100 else "Evaluasi record SLA 2 yang tidak memenuhi target."})
    units.sort(key=lambda x: (x['percentage'], x['denominator']), reverse=True)
    return {"sla": "SLA 2", "accumulation": {"total_records": total, "denominator": denominator, "on_time": on_time, "out_of_date": late, "incomplete": incomplete, "percentage": pct, "ontime_pct": pct, "ood_pct": round(late / denominator * 100, 2) if denominator else 0.0, "inc_pct": round(incomplete / total * 100, 2) if total else 0.0}, "unit_summary": units, "ranking": {"top5": units[:5], "bottom5": sorted(units, key=lambda x: (x['percentage'], -x['denominator']))[:5]}}


@router.get('/monitoring/sla2/traceability')
def sla2_traceability(page: int = 1, limit: int = 150, start_date: Optional[str] = None, end_date: Optional[str] = None, unit: Optional[str] = None, status: Optional[str] = None, db: Session = Depends(get_db)):
    q = db.query(SLARecord).filter(SLARecord.sla_code == SLA_CODE)
    if start_date:
        q = q.filter(SLARecord.registration_timestamp >= start_date)
    if end_date:
        q = q.filter(SLARecord.registration_timestamp <= end_date + " 23:59:59")
    if unit:
        q = q.filter(SLARecord.unit.ilike(f"%{unit}%"))
    q = _sla2_status_filter(q, status)
    total = q.count()
    records = q.order_by(SLARecord.id.desc()).offset((page - 1) * limit).limit(limit).all()
    units = sorted([x[0] for x in db.query(SLARecord.unit).filter(SLARecord.sla_code == SLA_CODE).distinct().all() if x[0]])
    return {"total": total, "page": page, "limit": limit, "units": units, "data": records}


@router.delete('/monitoring/{record_id}')
def delete_record(record_id: int, db: Session = Depends(get_db)):
    record = db.query(SLARecord).filter(SLARecord.id == record_id, SLARecord.sla_code == SLA_CODE).first()
    if not record:
        raise HTTPException(status_code=404, detail="Record tidak ditemukan")
    db.delete(record); db.commit()
    return {"message": "Record dihapus"}


@router.get('/export')
def export_excel(start_date: Optional[str] = None, end_date: Optional[str] = None, unit: Optional[str] = None, status: Optional[str] = None, db: Session = Depends(get_db)):
    q = _apply_filters(db.query(SLARecord).filter(SLARecord.sla_code == SLA_CODE), start_date, end_date, unit, status)
    rows = q.order_by(SLARecord.id).all()
    data = []
    for r in rows:
        data.append({
            "Nomor Registrasi": r.nomor_registrasi, "Nama Dokumen / Hal": r.nama_dokumen, "Unit": r.unit, "Sumber": r.sumber,
            "Status Registrasi": r.status_registrasi, "Status Inisiasi": r.status_inisiasi, "Posisi Data": r.posisi_data,
            "Format Arsip": r.format_arsip, "Document Type": r.document_type,
            "SLA 1A": r.sla1a_result, "SLA 1B": r.sla1b_result, "SLA 1 Final": r.status,
            "SLA 2A": r.sla2a_result, "SLA 2B": r.sla2b_result, "SLA 2 Final": r.sla2_final_result,
            "SLA 2 Manual": r.sla2_manual_result, "SLA 2 Decision Source": r.sla2_decision_source, "SLA 2 Manual Rule": r.sla2_manual_rule,
            "SLA 2A Working Days": r.sla2a_working_days, "SLA 2B Working Days": r.sla2b_working_days, "SLA 2 Working Days": r.sla2_working_days,
            "SLA 2 Flow Stage": r.sla2_flow_stage, "SLA 2 Reason": r.sla2_reason,
            "Flow Stage": r.flow_stage,
            "Decision Source": r.sla1_decision_source, "Manual Rule": r.sla1_manual_rule, "Working Days": r.sla1_working_days,
            "Reason": r.incomplete_reason, "Holiday Calendar Version": r.holiday_calendar_version, "Reporting Period": r.reporting_period,
        })
    out = io.BytesIO()
    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        pd.DataFrame(data).to_excel(writer, index=False, sheet_name="SLA 1 + SLA 2 Result")
    out.seek(0)
    return Response(out.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": 'attachment; filename="Monitoring_SLA_1_SLA_2_Result.xlsx"'})


# ------------------------- Holiday Calendar -------------------------

@router.get('/holidays')
def list_holidays(year: int = 2026, db: Session = Depends(get_db)):
    version = _latest_version(db, year)
    rows = db.query(HolidayCalendar).filter(HolidayCalendar.year == year).order_by(HolidayCalendar.holiday_date).all()
    return {"year": year, "version": version, "count": len(rows), "items": [{"id": r.id, "date": r.holiday_date.isoformat(), "description": r.description or "", "version": r.version} for r in rows]}


@router.post('/holidays')
def add_holiday(payload: dict, db: Session = Depends(get_db)):
    d = datetime.strptime(payload["date"], "%Y-%m-%d").date()
    year = d.year
    current_version = _latest_version(db, year)
    current = db.query(HolidayCalendar).filter(HolidayCalendar.year == year, HolidayCalendar.version == current_version).order_by(HolidayCalendar.holiday_date).all() if current_version else []
    if any(x.holiday_date == d for x in current):
        raise HTTPException(status_code=409, detail="Tanggal tersebut sudah ada di kalender aktif.")
    version = _next_version(db, year)
    for x in current:
        db.add(HolidayCalendar(year=year, holiday_date=x.holiday_date, description=x.description, version=version))
    db.add(HolidayCalendar(year=year, holiday_date=d, description=payload.get("description", ""), version=version))
    db.commit()
    return {"message": "Hari libur ditambahkan sebagai versi kalender baru.", "version": version}


@router.delete('/holidays/{holiday_id}')
def delete_holiday(holiday_id: int, db: Session = Depends(get_db)):
    row = db.query(HolidayCalendar).filter(HolidayCalendar.id == holiday_id).first()
    if not row:
        raise HTTPException(status_code=404, detail="Hari libur tidak ditemukan")
    year = row.year
    current_version = _latest_version(db, year)
    current = db.query(HolidayCalendar).filter(HolidayCalendar.year == year, HolidayCalendar.version == current_version).order_by(HolidayCalendar.holiday_date).all()
    version = _next_version(db, year)
    for x in current:
        if x.id != holiday_id:
            db.add(HolidayCalendar(year=year, holiday_date=x.holiday_date, description=x.description, version=version))
    db.commit()
    return {"message": "Hari libur dihapus dari kalender aktif.", "version": version}


@router.post('/holidays/import')
async def import_holidays(file: UploadFile = File(...), year: int = 2026, db: Session = Depends(get_db)):
    content = await file.read()
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        values = list(ws.iter_rows(values_only=True))
        if not values:
            raise ValueError("Excel kosong")
        header = [str(v or "").strip().lower() for v in values[0]]
        date_idx = next((i for i, h in enumerate(header) if h in {"tanggal", "date", "holiday", "hari libur", "holiday date"}), None)
        desc_idx = next((i for i, h in enumerate(header) if h in {"keterangan", "description", "holiday name", "nama hari libur"}), None)
        parsed = []
        if date_idx is not None:
            for row in values[1:]:
                if date_idx >= len(row):
                    continue
                d = _parse_date_any(row[date_idx])
                if d and d.year == year:
                    parsed.append((d, str(row[desc_idx] or "") if desc_idx is not None and desc_idx < len(row) else ""))
        else:
            ref = wb["99. Reference List"] if "99. Reference List" in wb.sheetnames else ws
            for r in range(3, 19):
                d = _parse_date_any(ref[f"C{r}"].value)
                if d and d.year == year:
                    parsed.append((d, "Reference List"))
        wb.close()
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Gagal membaca daftar hari libur: {e}")
    parsed = list(dict.fromkeys(parsed))
    if not parsed:
        raise HTTPException(status_code=400, detail="Tidak menemukan tanggal hari libur pada file.")
    current_version = _latest_version(db, year)
    current = db.query(HolidayCalendar).filter(HolidayCalendar.year == year, HolidayCalendar.version == current_version).all() if current_version else []
    merged = {(x.holiday_date, x.description or "") for x in current}
    merged.update(parsed)
    version = _next_version(db, year)
    for d, desc in sorted(merged):
        db.add(HolidayCalendar(year=year, holiday_date=d, description=desc, version=version))
    db.commit()
    return {"message": f"{len(parsed)} tanggal dari file diproses dan kalender aktif dibuat ulang.", "year": year, "version": version, "count": len(merged)}


def _next_version(db: Session, year: int) -> str:
    rows = db.query(HolidayCalendar.version).filter(HolidayCalendar.year == year).all()
    nums = []
    for (v,) in rows:
        try:
            nums.append(int(str(v).rsplit('-v', 1)[1]))
        except Exception:
            pass
    return f"HC-{year}-v{max(nums or [0]) + 1}"


def _parse_date_any(v):
    if isinstance(v, datetime): return v.date()
    if isinstance(v, date): return v
    if v is None or str(v).strip() == "": return None
    text = str(v).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d %b %Y", "%d %B %Y", "%m/%d/%Y"):
        try: return datetime.strptime(text, fmt).date()
        except ValueError: pass
    return None


@router.get('/holidays/preview')
def preview_workdays(start_date: str, end_date: str, year: Optional[int] = None, db: Session = Depends(get_db)):
    start = datetime.strptime(start_date, "%Y-%m-%d").date(); end = datetime.strptime(end_date, "%Y-%m-%d").date()
    if start > end: raise HTTPException(status_code=400, detail="START tidak boleh lebih besar dari END")
    years = range(start.year, end.year + 1)
    holiday_set = set()
    versions = {}
    for y in years:
        v = _latest_version(db, y); versions[y] = v
        holiday_set.update(_active_holidays(db, y, v))
    items = []
    cur = start
    while cur <= end:
        if cur.weekday() >= 5: reason = "Weekend"
        elif cur in holiday_set: reason = "Hari libur"
        else: reason = "Hari kerja"
        items.append({"date": cur.isoformat(), "is_working_day": reason == "Hari kerja", "reason": reason})
        cur += timedelta(days=1)
    networkdays = sum(1 for x in items if x["is_working_day"])
    return {"start": start_date, "end": end_date, "networkdays": networkdays, "networkdays_minus_one": networkdays - 1, "versions": versions, "days": items}
