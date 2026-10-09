from __future__ import annotations

from datetime import date, datetime, time, timedelta
from io import BytesIO
from typing import Any

from openpyxl import load_workbook


REQUIRED_HEADERS = {
    "sumber": [("sumber", 1)],
    "status_registrasi": [("status registrasi", 1)],
    "status_inisiasi": [("status inisiasi", 1)],
    "posisi_data": [("posisi data", 1)],
    "format_arsip": [("tipe arsip", 1), ("format arsip", 1)],
    "document_type": [("jenis dokumen", 1), ("document type", 1)],
    "tanggal_registrasi": [("tanggal registrasi", 2), ("tanggal registrasi", 1)],
    "jam_registrasi": [("jam registrasi", 1)],
    "tanggal_verifikasi_arsip": [("tanggal verifikasi arsip", 1)],
    "jam_verifikasi_arsip": [("jam verifikasi arsip", 1)],
    "tanggal_jadwal_penjemputan": [("tanggal jadwal penjemputan", 1)],
    "jam_jadwal_penjemputan": [("jam jadwal penjemputan", 1)],
    "tanggal_penjemputan_dokumen": [("tanggal penjemputan dokumen", 1)],
    "jam_penjemputan_dokumen": [("jam penjemputan dokumen", 1)],
}

OPTIONAL_HEADERS = {
    "nomor_registrasi": [("no. registrasi", 1), ("nomor registrasi", 1)],
    "nama_dokumen": [("hal. arsip", 1), ("nama dokumen / hal", 1), ("nama dokumen", 1)],
    "unit_sub_bidang": [("unit - sub bidang", 1)],
    "unit_umum": [("unit umum", 1)],
    "tanggal_penjadwalan_inaktif": [("tanggal penjadwalan arsip inaktif", 1)],
    "jam_penjadwalan_inaktif": [("jam penjadwalan arsip inaktif", 1)],
    "tanggal_penolakan_jadwal_penjemputan": [("tanggal penolakan jadwal penjemputan", 1)],
    "jam_penolakan_jadwal_penjemputan": [("jam penolakan jadwal penjemputan", 1)],
    "tanggal_gagal_penjemputan": [("tanggal gagal penjemputan", 1)],
    "jam_gagal_penjemputan": [("jam gagal penjemputan", 1)],
    "tanggal_penjadwalan_inaktif_2": [("tanggal penjadwalan arsip inaktif 2", 1)],
    "jam_penjadwalan_inaktif_2": [("jam penjadwalan arsip inaktif 2", 1)],
    "tanggal_runner_record_center": [
        ("tanggal runner sampai di record center", 1),
        ("tanggal runner sampai di rc", 1),
    ],
    "jam_runner_record_center": [
        ("jam runner sampai di record center", 1),
        ("jam runner sampai di rc", 1),
    ],
}


def norm_header(value: Any) -> str:
    return " ".join(str(value or "").strip().lower().replace("\n", " ").split())


def raw(value: Any) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, time):
        return value.isoformat()
    return str(value)


def as_date(value: Any) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d %b %Y", "%d %B %Y", "%d/%m/%Y", "%d/%m/%y", "%m/%d/%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    try:
        return datetime.fromisoformat(text).date()
    except Exception:
        return None


def as_time(value: Any) -> time | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.time()
    if isinstance(value, time):
        return value
    if isinstance(value, (float, int)) and 0 <= float(value) < 1:
        return (datetime.min + timedelta(seconds=round(float(value) * 86400))).time()
    text = str(value).strip()
    for fmt in ("%H:%M:%S", "%H:%M", "%I:%M:%S %p", "%I:%M %p"):
        try:
            return datetime.strptime(text, fmt).time()
        except ValueError:
            pass
    return None


def _find_index(headers: list[str], options: list[tuple[str, int]]) -> int | None:
    for needle, occurrence in options:
        matches = [idx for idx, header in enumerate(headers) if header == needle]
        if len(matches) >= occurrence:
            return matches[occurrence - 1]
    return None


def _score_headers(headers: list[str]) -> tuple[int, dict[str, int]]:
    mapping: dict[str, int] = {}
    for key, options in {**REQUIRED_HEADERS, **OPTIONAL_HEADERS}.items():
        idx = _find_index(headers, options)
        if idx is not None:
            mapping[key] = idx
    score = sum(1 for key in REQUIRED_HEADERS if key in mapping)
    return score, mapping


def _choose_unit(row: tuple[Any, ...], mapping: dict[str, int]) -> str | None:
    sub = row[mapping["unit_sub_bidang"]] if mapping.get("unit_sub_bidang") is not None else None
    umum = row[mapping["unit_umum"]] if mapping.get("unit_umum") is not None else None
    return raw(sub) or raw(umum)


def parse_sla4_workbook(file_bytes: bytes) -> tuple[list[dict[str, Any]], str, dict[str, Any]]:
    wb = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    try:
        best = None
        for ws in wb.worksheets:
            for header_row in range(1, min(ws.max_row, 40) + 1):
                row = list(next(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True)))
                headers = [norm_header(v) for v in row]
                score, mapping = _score_headers(headers)
                if best is None or score > best[0]:
                    best = (score, ws, header_row, mapping)
        if best is None or best[0] < len(REQUIRED_HEADERS):
            raise ValueError("Header raw SLA 4 tidak lengkap. Pastikan nama header sesuai raw data.")
        _, ws, header_row, mapping = best

        def val(row: tuple[Any, ...], key: str):
            idx = mapping.get(key)
            return row[idx] if idx is not None and idx < len(row) else None

        records: list[dict[str, Any]] = []
        registration_dates: list[date] = []
        for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
            no_reg = val(row, "nomor_registrasi")
            reg_date = as_date(val(row, "tanggal_registrasi"))
            name = val(row, "nama_dokumen")
            if all(v is None or str(v).strip() == "" for v in (no_reg, name, reg_date)):
                continue
            if reg_date:
                registration_dates.append(reg_date)

            item = {
                "nomor_registrasi": raw(no_reg),
                "nama_dokumen": raw(name),
                "unit": _choose_unit(row, mapping),
                "sumber": raw(val(row, "sumber")),
                "status_registrasi": raw(val(row, "status_registrasi")),
                "status_inisiasi": raw(val(row, "status_inisiasi")),
                "posisi_data": raw(val(row, "posisi_data")),
                "format_arsip": raw(val(row, "format_arsip")),
                "document_type": raw(val(row, "document_type")),
            }

            date_fields = [
                "tanggal_registrasi",
                "tanggal_verifikasi_arsip",
                "tanggal_penjadwalan_inaktif",
                "tanggal_jadwal_penjemputan",
                "tanggal_penolakan_jadwal_penjemputan",
                "tanggal_gagal_penjemputan",
                "tanggal_penjadwalan_inaktif_2",
                "tanggal_penjemputan_dokumen",
                "tanggal_runner_record_center",
            ]
            time_fields = [
                "jam_registrasi",
                "jam_verifikasi_arsip",
                "jam_penjadwalan_inaktif",
                "jam_jadwal_penjemputan",
                "jam_penolakan_jadwal_penjemputan",
                "jam_gagal_penjemputan",
                "jam_penjadwalan_inaktif_2",
                "jam_penjemputan_dokumen",
                "jam_runner_record_center",
            ]
            for field in date_fields:
                item[field] = as_date(val(row, field))
            for field in time_fields:
                item[field] = as_time(val(row, field))
            records.append(item)

        if not registration_dates:
            raise ValueError("Tidak ditemukan Tanggal Registrasi yang valid pada raw SLA 4.")
        period = max(registration_dates).strftime("%Y-%m")
        return records, period, {
            "sheet": ws.title,
            "header_row": header_row,
            "header_mapping": {key: idx + 1 for key, idx in mapping.items()},
            "record_count": len(records),
        }
    finally:
        wb.close()
