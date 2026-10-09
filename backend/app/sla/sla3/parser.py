from __future__ import annotations
from datetime import date, datetime, time
from io import BytesIO
from typing import Any
from openpyxl import load_workbook

REQUIRED_HEADERS = {
    "sumber": [("sumber", 1)],
    "status_registrasi": [("status registrasi", 1)],
    "status_inisiasi": [("status inisiasi", 1)],
    "posisi_data": [("posisi data", 1)],
    "format_arsip": [("tipe arsip", 1), ("format arsip", 1)],
    "document_type": [("jenis dokumen", 1), ("document type", 1)],
    "tanggal_registrasi": [("tanggal registrasi", 2), ("tanggal registrasi", 1)],
}

def parse_sla3_excel(file_bytes: bytes, filename: str) -> list[dict[str, Any]]:
    wb = load_workbook(BytesIO(file_bytes), data_only=True)
    sheet = wb.active
    
    # Deteksi baris header
    header_row_idx = 1
    col_map = {}
    for r in range(1, min(10, sheet.max_row + 1)):
        row_vals = [str(sheet.cell(row=r, column=c).value or "").strip().lower() for c in range(1, sheet.max_column + 1)]
        if any("tanggal registrasi" in v for v in row_vals) or any("status registrasi" in v for v in row_vals):
            header_row_idx = r
            for c_idx, val in enumerate(row_vals, start=1):
                if val:
                    col_map[val] = c_idx
            break

    if not col_map:
        for c in range(1, sheet.max_column + 1):
            val = str(sheet.cell(row=1, column=c).value or "").strip().lower()
            if val:
                col_map[val] = c

    def get_cell_val(row_idx, names):
        for name in names:
            c = col_map.get(name)
            if c:
                return sheet.cell(row=row_idx, column=c).value
        return None

    def parse_date(val):
        if isinstance(val, datetime):
            return val.date()
        if isinstance(val, date):
            return val
        if val:
            try:
                return datetime.strptime(str(val).strip()[:10], "%Y-%m-%d").date()
            except ValueError:
                try:
                    return datetime.strptime(str(val).strip()[:10], "%d/%m/%Y").date()
                except Exception:
                    pass
        return None

    records = []
    for r in range(header_row_idx + 1, sheet.max_row + 1):
        reg_date = parse_date(get_cell_val(r, ["tanggal registrasi"]))
        if not reg_date:
            # Cek apakah baris kosong total
            row_empty = all(sheet.cell(row=r, column=c).value is None for c in range(1, sheet.max_column + 1))
            if row_empty:
                continue

        record = {
            "nomor_registrasi": str(get_cell_val(r, ["no. registrasi", "nomor registrasi"]) or "").strip(),
            "nama_dokumen": str(get_cell_val(r, ["hal. arsip", "nama dokumen"]) or "").strip(),
            "unit": str(get_cell_val(r, ["unit", "unit - sub bidang"]) or "Tanpa Unit").strip(),
            "sumber": str(get_cell_val(r, ["sumber"]) or "").strip(),
            "status_registrasi": str(get_cell_val(r, ["status registrasi"]) or "").strip(),
            "status_inisiasi": str(get_cell_val(r, ["status inisiasi"]) or "").strip(),
            "posisi_data": str(get_cell_val(r, ["posisi data"]) or "").strip(),
            "format_arsip": str(get_cell_val(r, ["tipe arsip", "format arsip"]) or "").strip(),
            "document_type": str(get_cell_val(r, ["jenis dokumen", "document type"]) or "").strip(),
            "tanggal_registrasi": reg_date,
            "tanggal_penjadwalan_inaktif": parse_date(get_cell_val(r, ["tanggal penjadwalan arsip inaktif"])),
            "tanggal_jadwal_penjemputan": parse_date(get_cell_val(r, ["tanggal jadwal penjemputan"])),
            "tanggal_penjadwalan_inaktif_2": parse_date(get_cell_val(r, ["tanggal penjadwalan arsip inaktif 2"])),
        }
        records.append(record)
    return records