from datetime import date, datetime
from io import BytesIO

from openpyxl import Workbook, load_workbook

from app.services.excel_parser import parse_registrasi_sheet
from app.sla.sla4.calculator import SLA4Calculator
from app.sla.sla4.unit_area_service import DEFAULT_VENDOR_UNITS, normalize_unit


def test_shared_parser_keeps_sla4_milestone_fields():
    wb = Workbook()
    ws = wb.active
    ws.append([
        "Sumber", "Status Registrasi", "Status Inisiasi", "Posisi Data", "Nomor Registrasi",
        "Nama Dokumen / Hal", "Unit", "Format Arsip", "Document Type", "Tanggal Registrasi",
        "Jam Registrasi", "Tanggal Verifikasi Arsip", "Jam Verifikasi Arsip",
        "Tanggal Penjadwalan Arsip Inaktif", "Jam Penjadwalan Arsip Inaktif",
        "Tanggal Penjadwalan Arsip Inaktif 2", "Jam Penjadwalan Arsip Inaktif 2",
        "Tanggal Penolakan Jadwal Penjemputan", "Jam Penolakan Jadwal Penjemputan",
        "Tanggal Gagal Penjemputan", "Jam Gagal Penjemputan",
        "Tanggal Jadwal Penjemputan", "Jam Jadwal Penjemputan",
        "Tanggal Penjemputan Dokumen", "Jam Penjemputan Dokumen",
        "Tanggal Runner sampai di Record Center", "Jam Runner sampai di Record Center",
    ])
    ws.append([
        "ELARCH", "Diarsipkan", "-", "Record Center", "REG-001", "Dokumen", "Kantor Pusat",
        "Fisik", "Inaktif", "06 Aug 2026", "08:00:00", "06 Aug 2026", "08:10:00",
        "20 Aug 2026", "09:00:00", "21 Aug 2026", "09:00:00", None, None, None, None,
        "20 Aug 2026", "09:00:00", "21 Aug 2026", "10:00:00", "21 Aug 2026", "11:00:00"
    ])
    out = BytesIO(); wb.save(out)
    records, context = parse_registrasi_sheet(out.getvalue(), holidays=[])
    row = records[0]
    assert context["period_label"] == "2026-08"
    assert row["tanggal_jadwal_penjemputan"] == date(2026, 8, 20)
    assert row["tanggal_penjemputan_dokumen"] == date(2026, 8, 21)
    assert row["tanggal_runner_record_center"] == date(2026, 8, 21)


def test_sla4_expected_august_counts_with_shared_parser():
    source = "/mnt/data/PLN_Archive_Service_Hub_SLA_Calculations_August_2026_vFinal_080920261.xlsx"
    with open(source, "rb") as f:
        records, context = parse_registrasi_sheet(f.read(), holidays=[])

    wb = load_workbook(source, read_only=True, data_only=True)
    ref = wb["99. Reference List"]
    holidays = []
    for r in range(3, 21):
        value = ref[f"C{r}"].value
        if hasattr(value, "date"):
            value = value.date()
        if isinstance(value, date):
            holidays.append(value)
    wb.close()

    unit_map = {normalize_unit(name): area for name, _unit_type, area in DEFAULT_VENDOR_UNITS}
    final = {"Yes": 0, "No": 0, "N/A": 0}
    for row in records:
        result = SLA4Calculator.evaluate(row, holidays, unit_map.get(normalize_unit(row.get("unit"))))
        final[result["sla4_final_result"]] += 1

    assert context["period_label"] == "2026-08"
    assert final == {"Yes": 6684, "No": 0, "N/A": 22302}
